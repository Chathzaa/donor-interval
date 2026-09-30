"""Generate the editable three-minute EC8204 slide template."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


OUT = Path(__file__).with_name("GP_XX_DonorInterval_Presentation.pptx")
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

BG = RGBColor(9, 19, 32)
PANEL = RGBColor(18, 40, 52)
PANEL2 = RGBColor(26, 54, 63)
MINT = RGBColor(78, 225, 176)
WHITE = RGBColor(239, 249, 250)
MUTED = RGBColor(170, 195, 205)
AMBER = RGBColor(244, 189, 111)


def box(slide, x, y, w, h, fill=PANEL, radius=True):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    return shape


def txt(slide, value, x, y, w, h, size=20, color=WHITE, bold=False,
        align=PP_ALIGN.LEFT, font="Aptos"):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(0.02)
    frame.margin_top = Inches(0.01)
    para = frame.paragraphs[0]
    para.text = value
    para.alignment = align
    para.font.name = font
    para.font.size = Pt(size)
    para.font.bold = bold
    para.font.color.rgb = color
    return shape


def slide_base(number, label):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = BG
    txt(slide, "DONORINTERVAL  /  EC8204", .65, .33, 6, .35, 11, MINT, True)
    txt(slide, label.upper(), 9.6, .33, 3.1, .35, 10, MUTED, True, PP_ALIGN.RIGHT)
    box(slide, .66, 7.04, 12.0, .012, PANEL2, False)
    txt(slide, f"GROUP [XX]   •   EDUCATIONAL PROTOTYPE", .7, 7.1, 7, .22, 9, MUTED)
    txt(slide, f"{number:02d} / 05", 11.55, 7.1, 1.0, .22, 9, MUTED, align=PP_ALIGN.RIGHT)
    return slide


# Slide 1 — problem and promise.
s = slide_base(1, "The problem")
txt(s, "One shared timeline.\nSafer donation scheduling.", .8, 1.3, 9.5, 1.75, 36, WHITE, True)
txt(s, "Blockchain-Based Blood Donation Interval Registry", .82, 3.47, 9.5, .48, 21, MINT, True)
txt(s, "Participating blood banks need to see recent recorded donations before adding a new one.", .82, 4.08, 9.3, .9, 19, MUTED)
box(s, .8, 5.38, 11.73, .95)
txt(s, "Group [XX]  •  [Member 1]  •  [Member 2]  •  [Member 3]", 1.1, 5.67, 11.1, .4, 17, WHITE, True)

# Slide 2 — architecture.
s = slide_base(2, "How it works")
txt(s, "Authorized banks write. Anyone with the code can check.", .75, .95, 11.8, .65, 26, WHITE, True)
cards = [
    ("BANK A", "Records a donation using a fictional donor code hash."),
    ("SMART CONTRACT", "Checks the interval, stores a timestamp, emits an audit event."),
    ("BANK B", "Reads the same history and cannot record an early second donation."),
]
for i, (title, body) in enumerate(cards):
    x = .75 + i * 4.18
    box(s, x, 2.06, 3.8, 2.25)
    txt(s, title, x + .28, 2.39, 3.25, .4, 16, MINT, True)
    txt(s, body, x + .28, 2.93, 3.2, 1.1, 18, WHITE)
txt(s, "→", 4.53, 2.8, .35, .5, 30, MINT, True)
txt(s, "←", 8.7, 2.8, .35, .5, 30, MINT, True)
box(s, .75, 4.85, 12.15, 1.12, PANEL2)
txt(s, "OFF-CHAIN: real identity, blood group, screening, and the donor-code mapping", 1.05, 5.16, 11.5, .5, 18, WHITE, True)

# Slide 3 — rules and demonstration.
s = slide_base(3, "Live demonstration")
txt(s, "One code. Two banks. A rejected early record.", .75, .94, 11.9, .65, 27, WHITE, True)
steps = [
    ("01", "Generate", "Use a random fictional donor code."),
    ("02", "Record", "Bank A writes the first timestamp."),
    ("03", "Check", "Bank B sees the shared history."),
    ("04", "Reject", "The contract blocks another record too soon."),
]
for i, (num, title, body) in enumerate(steps):
    y = 1.91 + i * 1.04
    box(s, .75, y, 5.1, .85, PANEL)
    txt(s, num, 1.0, y + .13, .5, .3, 15, MINT, True)
    txt(s, title, 1.61, y + .09, 3.95, .33, 17, WHITE, True)
    txt(s, body, 1.62, y + .43, 4.05, .32, 11, MUTED)
box(s, 6.15, 1.91, 6.44, 4.2, PANEL2)
s.shapes.add_picture(
    str(Path(__file__).with_name("app-preview.png")),
    Inches(6.3), Inches(2.06), width=Inches(6.15), height=Inches(3.9)
)
txt(s, "124-day fixed interval is a conservative demonstration approximation, not a clinical decision.", .82, 6.33, 11.8, .42, 13, AMBER)

# Slide 4 — evidence.
s = slide_base(4, "Implementation & tests")
txt(s, "Tested security behavior", .75, .95, 11.5, .62, 28, WHITE, True)
box(s, .75, 1.92, 3.45, 3.8, PANEL2)
txt(s, "8", 1.1, 2.28, 2.7, 1.18, 72, MINT, True)
txt(s, "automated tests passing", 1.12, 3.55, 2.6, 1.2, 20, WHITE, True)
cases = [
    "First donation and timestamp",
    "Early second donation blocked across banks",
    "Later donation accepted after simulated time",
    "Unauthorized or revoked bank blocked",
    "Independent donor histories and invalid inputs",
]
for i, item in enumerate(cases):
    y = 1.95 + i * .79
    box(s, 4.55, y, 8.05, .62, PANEL)
    txt(s, "✓", 4.86, y + .12, .42, .3, 17, MINT, True)
    txt(s, item, 5.45, y + .13, 6.9, .33, 16, WHITE)
txt(s, "Solidity smart contract  •  Hardhat local chain  •  MetaMask + ethers web app", .8, 6.24, 11.8, .45, 15, MUTED)

# Slide 5 — honest conclusions.
s = slide_base(5, "Security & limits")
txt(s, "What the registry proves — and what it cannot", .75, .95, 11.7, .7, 28, WHITE, True)
box(s, .75, 1.9, 5.75, 3.55)
txt(s, "PROVES", 1.12, 2.25, 4.9, .4, 17, MINT, True)
txt(s, "Which authorized wallet submitted a pseudonymous record, when it was submitted, and whether the recorded interval has passed.", 1.12, 2.86, 4.87, 1.65, 21, WHITE)
box(s, 6.76, 1.9, 5.82, 3.55)
txt(s, "DOES NOT PROVE", 7.14, 2.25, 4.95, .4, 17, AMBER, True)
txt(s, "A donor's medical eligibility, real identity, or donations outside participating banks. A disclosed code also links public records.", 7.14, 2.86, 4.9, 1.65, 21, WHITE)
txt(s, "A shared audit trail supports coordination; clinical staff retain the final decision.", .82, 5.95, 11.75, .65, 19, MINT, True)

prs.save(OUT)
print(OUT)
